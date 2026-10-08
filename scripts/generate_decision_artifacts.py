#!/usr/bin/env python3
"""Generate the decision-analysis Excel workbook and PowerPoint deck.

Single source of truth: the criteria, weights and alternative scores defined
in DATA below. Edit them and re-run the script to regenerate both artifacts.

Usage (from the repository root):

    python3 -m venv .venv && . .venv/bin/activate
    pip install openpyxl python-pptx
    python scripts/generate_decision_artifacts.py [--out-dir docs/decision-model]

Outputs:
    <out-dir>/decision-model.xlsx
    <out-dir>/decision-model.pptx
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Source data (edit here only)
# ---------------------------------------------------------------------------

PROBLEM = {
    "المشكلة": "",
    "القرار المطلوب": "",
    "صاحب القرار": "",
    "تاريخ القرار": "",
    "القيود": "",
    "الافتراضات": "",
}

# (criterion, weight, note)
CRITERIA = [
    ("التكلفة", 0.30, ""),
    ("السرعة", 0.20, ""),
    ("المخاطر", 0.25, ""),
    ("الأثر الاستراتيجي", 0.25, ""),
]

ALTERNATIVES = ["A", "B", "C"]

# SCORES[criterion_index][alternative_index]
SCORES = [
    [8, 7, 9],  # التكلفة
    [7, 9, 6],  # السرعة
    [6, 8, 5],  # المخاطر
    [9, 7, 8],  # الأثر الاستراتيجي
]

# Implementation plan rows: (activity, owner, date, resources, KPI, status)
PLAN = [
    ("", "", "", "", "", ""),
    ("", "", "", "", "", ""),
]

EXCEL_NAME = "decision-model.xlsx"
PPTX_NAME = "decision-model.pptx"

# ---------------------------------------------------------------------------
# Computation (mirrors the Excel formulas, used for the deck and checks)
# ---------------------------------------------------------------------------


def validate() -> None:
    total = sum(w for _, w, _ in CRITERIA)
    if not math.isclose(total, 1.0, abs_tol=1e-9):
        raise ValueError(f"Criteria weights must sum to 1.00, got {total:.2f}")
    for row in SCORES:
        if len(row) != len(ALTERNATIVES):
            raise ValueError("Each criterion needs one score per alternative")


def weighted_scores() -> list[float]:
    return [
        round(sum(CRITERIA[i][1] * SCORES[i][j] for i in range(len(CRITERIA))), 2)
        for j in range(len(ALTERNATIVES))
    ]


def ranks(scores: list[float]) -> list[int]:
    ordered = sorted(scores, reverse=True)
    return [ordered.index(s) + 1 for s in scores]


# ---------------------------------------------------------------------------
# Excel workbook
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)


def _style_header(ws, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER


def _set_widths(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def build_workbook(path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    # Sheet 1: problem and decision
    ws1 = wb.create_sheet("المشكلة والقرار")
    ws1.sheet_view.rightToLeft = True
    ws1.append(["البند", "التفاصيل"])
    _style_header(ws1, 1, 2)
    for key, value in PROBLEM.items():
        ws1.append([key, value])
    _set_widths(ws1, [24, 60])

    # Sheet 2: criteria and weights
    ws2 = wb.create_sheet("المعايير والأوزان")
    ws2.sheet_view.rightToLeft = True
    ws2.append(["المعيار", "الوزن", "ملاحظات"])
    _style_header(ws2, 1, 3)
    for name, weight, note in CRITERIA:
        ws2.append([name, weight, note])
    first, last = 2, 1 + len(CRITERIA)
    total_row = last + 1
    ws2.append(["المجموع", f"=SUM(B{first}:B{last})", ""])
    for cell in ws2[total_row]:
        cell.font = TOTAL_FONT
    for row in ws2.iter_rows(min_row=first, max_row=total_row, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "0.00"
            cell.alignment = CENTER
    _set_widths(ws2, [28, 12, 40])

    # Sheet 3: alternative evaluation (weights linked to sheet 2)
    ws3 = wb.create_sheet("تقييم البدائل")
    ws3.sheet_view.rightToLeft = True
    headers = ["المعيار", "الوزن"] + [f"البديل {a}" for a in ALTERNATIVES]
    ws3.append(headers)
    _style_header(ws3, 1, len(headers))
    for i, (name, _, _) in enumerate(CRITERIA):
        r = first + i
        row = [name, f"='المعايير والأوزان'!B{r}"] + [
            SCORES[i][j] for j in range(len(ALTERNATIVES))
        ]
        ws3.append(row)
    eval_last = first + len(CRITERIA) - 1
    eval_total = eval_last + 1
    total = ["الدرجة المرجحة", ""]
    for j in range(len(ALTERNATIVES)):
        col = get_column_letter(3 + j)
        total.append(f"=SUMPRODUCT($B${first}:$B${eval_last},{col}{first}:{col}{eval_last})")
    ws3.append(total)
    for cell in ws3[eval_total]:
        cell.font = TOTAL_FONT
    for row in ws3.iter_rows(min_row=first, max_row=eval_total, min_col=2, max_col=2 + len(ALTERNATIVES)):
        for cell in row:
            cell.alignment = CENTER
    for row in ws3.iter_rows(min_row=first, max_row=eval_last, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "0.00"
    for row in ws3.iter_rows(min_row=eval_total, max_row=eval_total, min_col=3, max_col=2 + len(ALTERNATIVES)):
        for cell in row:
            cell.number_format = "0.00"
    _set_widths(ws3, [24, 10] + [14] * len(ALTERNATIVES))

    # Sheet 4: decision matrix (linked to sheet 3)
    ws4 = wb.create_sheet("مصفوفة القرار")
    ws4.sheet_view.rightToLeft = True
    ws4.append(["البديل", "الدرجة", "الترتيب", "القرار"])
    _style_header(ws4, 1, 4)
    for j, alt in enumerate(ALTERNATIVES):
        r = 2 + j
        col = get_column_letter(3 + j)
        ws4.append(
            [
                alt,
                f"='تقييم البدائل'!{col}{eval_total}",
                f"=RANK(B{r},$B$2:$B${1 + len(ALTERNATIVES)})",
                f'=IF(C{r}=1,"✅ مختار","")',
            ]
        )
    for row in ws4.iter_rows(min_row=2, max_row=1 + len(ALTERNATIVES), min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "0.00"
    for row in ws4.iter_rows(min_row=2, max_row=1 + len(ALTERNATIVES), min_col=1, max_col=4):
        for cell in row:
            cell.alignment = CENTER
    _set_widths(ws4, [12, 12, 12, 16])

    # Sheet 5: implementation and follow-up
    ws5 = wb.create_sheet("خطة التنفيذ والمتابعة")
    ws5.sheet_view.rightToLeft = True
    ws5.append(["النشاط", "المسؤول", "الموعد", "الموارد", "مؤشر النجاح", "الحالة"])
    _style_header(ws5, 1, 6)
    for row in PLAN:
        ws5.append(list(row))
    _set_widths(ws5, [28, 18, 14, 22, 24, 14])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------------------
# PowerPoint deck
# ---------------------------------------------------------------------------

TITLE_COLOR = RGBColor(0x1F, 0x4E, 0x78)


def _rtl(paragraph) -> None:
    """Make a paragraph right-to-left (Arabic)."""
    paragraph.alignment = PP_ALIGN.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")


def _fill_text(tf, lines: list[str], size: int = 20) -> None:
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        _rtl(p)


def _slide_with_title(prs: Presentation, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = title
    t = slide.shapes.title.text_frame.paragraphs[0]
    _rtl(t)
    t.font.size = Pt(32)
    t.font.bold = True
    t.font.color.rgb = TITLE_COLOR
    return slide


def _add_body(slide, lines: list[str], top: float = 1.7, height: float = 5.0, size: int = 20):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(12.0), Inches(height))
    _fill_text(box.text_frame, lines, size=size)
    return box


def build_deck(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    scores = weighted_scores()
    rnk = ranks(scores)
    best = ALTERNATIVES[scores.index(max(scores))]

    # 1. Cover
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = "تحليل البيانات واتخاذ القرار"
    _rtl(cover.shapes.title.text_frame.paragraphs[0])
    cover.placeholders[1].text = "من البيانات إلى القرار"
    _rtl(cover.placeholders[1].text_frame.paragraphs[0])

    # 2. Problem
    s = _slide_with_title(prs, "المشكلة")
    _add_body(s, ["[وصف مختصر للمشكلة]"])

    # 3. Decision required
    s = _slide_with_title(prs, "القرار المطلوب")
    _add_body(s, ["[ما الذي نريد حسمه؟ ومن صاحب القرار؟]"])

    # 4. Data and sources
    s = _slide_with_title(prs, "البيانات والمصادر")
    _add_body(s, ["[مصادر البيانات وحجمها وفترتها]"])

    # 5. Methodology
    s = _slide_with_title(prs, "منهجية التحليل")
    _add_body(
        s,
        [
            "• وصفي: ماذا حدث؟",
            "• تشخيصي: لماذا حدث؟",
            "• تنبؤي: ماذا سيحدث؟",
            "• توجيهي: ماذا نفعل؟",
        ],
    )

    # 6. Findings
    s = _slide_with_title(prs, "النتائج والرؤى")
    _add_body(s, ["[أهم 3 إلى 5 نتائج]"])

    # 7. Alternatives
    s = _slide_with_title(prs, "البدائل المطروحة")
    _add_body(s, [f"• البديل {a}" for a in ALTERNATIVES])

    # 8. Evaluation matrix (table)
    s = _slide_with_title(prs, "مصفوفة التقييم")
    rows = len(CRITERIA) + 2
    cols = 2 + len(ALTERNATIVES)
    table = s.shapes.add_table(rows, cols, Inches(0.7), Inches(1.8), Inches(12.0), Inches(0.6) * rows).table
    header = ["المعيار", "الوزن"] + [f"البديل {a}" for a in ALTERNATIVES]
    for c, text in enumerate(header):
        table.cell(0, c).text = text
    for r, (name, weight, _) in enumerate(CRITERIA, start=1):
        table.cell(r, 0).text = name
        table.cell(r, 1).text = f"{weight:.2f}"
        for j in range(len(ALTERNATIVES)):
            table.cell(r, 2 + j).text = str(SCORES[r - 1][j])
    table.cell(rows - 1, 0).text = "الدرجة المرجحة"
    table.cell(rows - 1, 1).text = ""
    for j, sc in enumerate(scores):
        table.cell(rows - 1, 2 + j).text = f"{sc:.2f}"
    for r in range(rows):
        for c in range(cols):
            for p in table.cell(r, c).text_frame.paragraphs:
                _rtl(p)
                p.font.size = Pt(18)

    # 9. Recommendation
    s = _slide_with_title(prs, "القرار والتوصية")
    _add_body(
        s,
        [
            f"التوصية: البديل {best} (الدرجة {max(scores):.2f})",
            "الترتيب: " + "، ".join(
                f"{ALTERNATIVES[j]} = {scores[j]:.2f} (المركز {rnk[j]})"
                for j in range(len(ALTERNATIVES))
            ),
            "[السبب ومخاطر القرار]",
        ],
    )

    # 10. Implementation and follow-up
    s = _slide_with_title(prs, "خطة التنفيذ والمتابعة")
    _add_body(s, ["[المسؤوليات والمواعيد ومؤشرات النجاح]"])

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out-dir", default="docs/decision-model", help="output directory")
    args = parser.parse_args()

    validate()
    out = Path(args.out_dir)
    build_workbook(out / EXCEL_NAME)
    build_deck(out / PPTX_NAME)

    scores = weighted_scores()
    print("Weighted scores:", dict(zip(ALTERNATIVES, scores)))
    print(f"Wrote {out / EXCEL_NAME}")
    print(f"Wrote {out / PPTX_NAME}")


if __name__ == "__main__":
    main()
